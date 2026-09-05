"""PBS Punjab census facts, kept separate from the fixed reviewed scoring cohort."""
from __future__ import annotations

import argparse
from dataclasses import asdict
from decimal import Decimal, InvalidOperation
from hashlib import sha256
import json
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert

from app.facts import format_fact
from app.ingest_official_data import CONFIG, ROOT, Report, extract, load_observations, parse_count, resolve_district

SNAPSHOT = ROOT / 'data/curated/punjab_census.json'
LIMITATION = ('Population and education facts are available from Census 2023. Reviewed GPIS scores, '
              'health facilities, budget figures and comparable trends are not yet available for this district.')


def parse_percent(value):
    if value is None or str(value).strip().lower() in {'', '-', '..', 'n/a'}:
        raise ValueError('Missing percentage')
    if isinstance(value, bool):
        raise ValueError('Invalid percentage')
    try:
        number = Decimal(str(value).strip().removesuffix('%'))
    except InvalidOperation:
        raise ValueError('Invalid percentage') from None
    if not number.is_finite() or not 0 <= number <= 100:
        raise ValueError('Percentage outside 0–100')
    return float(number.quantize(Decimal('0.000001')))


def extract_education(path: Path, config: dict, population: list[dict], districts: dict):
    if sha256(path.read_bytes()).hexdigest() != config['sha256']:
        raise ValueError('PBS Table 12 checksum mismatch')
    population_by_slug = {item['slug']: item for item in population}
    workbook = load_workbook(path, read_only=True, data_only=False)
    observations = []
    seen = set()
    seen_blocks = set()
    try:
        sheet = workbook[config['sheet']]
        if ('TABLE 12' not in str(sheet['A1'].value) or 'CENSUS-2023' not in str(sheet['A1'].value)
                or [sheet[f'{column}3'].value for column in 'ABCD'] != ['INDICATORS', 'TOTAL', 'MALE', 'FEMALE']
                or sheet['B2'].value != 'TOTAL'):
            raise ValueError('Unexpected PBS Table 12 headers')
        slug = None
        for number, row in enumerate(sheet.iter_rows(min_row=5, values_only=True), 5):
            label = str(row[0] or '').strip()
            # Every geography heading starts a new block; never inherit a district into a tehsil.
            if label and all(value is None for value in row[1:]):
                slug = resolve_district(label, districts) if label.endswith(' DISTRICT') else None
                if slug:
                    if slug in seen_blocks:
                        raise ValueError(f'Duplicate district block: {slug}')
                    seen_blocks.add(slug)
                continue
            if slug is None or label not in config['rows']:
                continue
            for column, metric in config['rows'][label].items():
                key = (slug, metric)
                if key in seen:
                    raise ValueError(f'Duplicate Table 12 metric: {key}')
                seen.add(key)
                original = row[ord(column) - ord('A')]
                value = parse_percent(original) if metric.endswith('_pct') else parse_count(original)
                if value is None:
                    raise ValueError(f'Missing {metric} for {slug}')
                pop = population_by_slug[slug]
                observations.append(dict(slug=slug, name=pop['name'], population_2023=pop['value'],
                    value=value, metric_key=metric, unit='percent' if metric.endswith('_pct') else 'count',
                    category='education', period='2023', source_id=config['source_id'],
                    locator=f"Table 12; {config['sheet']}!{column}{number}; {pop['name']} District; "
                            f"{label}={original!r}; extraction=openpyxl; parser={config['parser_version']}; "
                            f"sha256={config['sha256']}"))
    finally:
        workbook.close()
    expected = {(slug, metric) for slug in population_by_slug for mapping in config['rows'].values() for metric in mapping.values()}
    if seen != expected:
        raise ValueError(f'Incomplete Table 12 district metrics: {sorted(expected - seen)}')
    return observations


def build_snapshot():
    pilot = json.loads((ROOT / 'data/curated/districts.json').read_text())
    configs = [json.loads(CONFIG.read_text()), json.loads((ROOT / 'data/curated/pbs_table_12.json').read_text())]
    population, report = extract(ROOT / pilot['sources']['pbs_table_1']['local_file'], configs[0])
    if report.rejected or report.missing or len(population) != len(configs[0]['districts']):
        raise ValueError('Incomplete population coverage')
    education = extract_education(ROOT / pilot['sources']['pbs_table_12']['local_file'], configs[1], population, configs[0])
    all_observations = population + education
    by_key = {(item['slug'], item['metric_key']): item for item in all_observations}
    for district in pilot['districts']:
        for metric, value in {'population_2023': district['population_2023'], **district['education']}.items():
            if abs(by_key[district['slug'], metric]['value'] - value) > 0.000001:
                raise ValueError(f'Official source disagrees with reviewed pilot: {district["slug"]}/{metric}')
    pilot_slugs = {item['slug'] for item in pilot['districts']}
    records = []
    for pop in population:
        if pop['slug'] in pilot_slugs:
            continue
        observations = [item for item in all_observations if item['slug'] == pop['slug']]
        evidence = []
        for category, source_id in [('context', 'pbs_table_1'), ('education', 'pbs_table_12')]:
            items = [item for item in observations if item['source_id'] == source_id]
            evidence.append(dict(id=f"{pop['slug']}-census-{category}", category=category,
                text=' '.join(format_fact(pop['name'], item['metric_key'], item['value'], item['unit']) for item in items),
                source_id=source_id, locator='; '.join(item['locator'] for item in items)))
        records.append(dict(slug=pop['slug'], name=pop['name'], population_2023=pop['value'],
            education={item['metric_key']: item['value'] for item in observations if item['source_id'] == 'pbs_table_12'},
            health={}, accountability={}, evidence=evidence, observations=observations))
    return dict(sources={config['source_id']: {**pilot['sources'][config['source_id']], 'sha256': config['sha256']}
                         for config in configs}, districts=records)


def unscored_record(record):
    return {**record, 'derived': {'scores': {}, 'score_details': {}, 'facility_count': None, 'facilities_per_100k': None},
            'data_warnings': [LIMITATION]}


def load_snapshot(session, snapshot):
    from app.database.models import District, EvidenceChunk
    report = Report(accepted=sum(len(record['observations']) for record in snapshot['districts']))
    for source_id, source in snapshot['sources'].items():
        observations = [item for record in snapshot['districts'] for item in record['observations'] if item['source_id'] == source_id]
        load_observations(session, observations, source, {'source_id': source_id, 'period': '2023', 'sha256': source['sha256']}, report)
    if report.rejected:
        raise ValueError('Existing observations conflict; census load rolled back: ' + '; '.join(report.warnings))
    for record in snapshot['districts']:
        district_id = session.scalar(select(District.id).where(District.slug == record['slug']))
        for evidence in record['evidence']:
            session.execute(insert(EvidenceChunk).values(**evidence, district_id=district_id, period='2023')
                            .on_conflict_do_nothing(index_elements=['id']))
    report.districts_covered = sorted(record["slug"] for record in snapshot["districts"])
    report.metrics_extracted = sorted({item["metric_key"] for record in snapshot["districts"] for item in record["observations"]})
    return report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--dry-run', action='store_true')
    args = parser.parse_args()
    snapshot = build_snapshot()
    if args.dry_run:
        print(json.dumps({'districts_total': 36, 'additional_districts': len(snapshot['districts']),
                          'additional_observations': sum(len(item['observations']) for item in snapshot['districts'])}))
        return
    from app.config import Settings
    from app.database.models import DatasetMetadata
    from app.database.session import session_factory
    settings = Settings.from_env()
    if not settings.database_url:
        parser.error('GPIS_DATABASE_URL is required')
    # Commit the verified source snapshot for reproducible JSON and database ingestion.
    output = json.dumps(snapshot, indent=2) + '\n'
    with session_factory(settings.database_url).begin() as session:
        report = load_snapshot(session, snapshot)
        from app.cache import retrieval_corpus_version
        version = sha256((retrieval_corpus_version(settings.curated_data_path, settings.score_version) + output).encode()).hexdigest()
        session.execute(insert(DatasetMetadata).values(key='retrieval_corpus_version', value=version)
                        .on_conflict_do_update(index_elements=['key'], set_={'value': version}))
    SNAPSHOT.write_text(output, encoding='utf-8')
    print(json.dumps(asdict(report), indent=2))


if __name__ == '__main__':
    main()
