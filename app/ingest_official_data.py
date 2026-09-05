"""Deterministic ingestion of the registered PBS Census 2023 Punjab Table 1."""
from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import asdict, dataclass, field
from decimal import Decimal, InvalidOperation
from hashlib import sha256
import json
from pathlib import Path
import re

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.dialects.postgresql import insert

from app.database.models import District, Observation, Source
from app.database.session import session_factory

ROOT = Path(__file__).resolve().parents[1]
CONFIG = ROOT / 'data/curated/pbs_table_1.json'


def district_key(value: str) -> str:
    return re.sub(r'[^a-z0-9]', '', re.sub(r'\s+district\s*$', '', value, flags=re.I).lower())


def resolve_district(value: str, config: dict) -> str:
    aliases = {district_key(name): slug for slug, name in config['districts'].items()}
    aliases.update({district_key(name): slug for name, slug in config['aliases'].items()})
    try:
        return aliases[district_key(value)]
    except KeyError:
        raise ValueError(f'Unknown district: {value}') from None


def parse_count(value) -> int | None:
    if value is None or str(value).strip().lower() in {'', '-', '..', '...', 'n/a', 'na'}:
        return None
    text = str(value).strip()
    if isinstance(value, bool) or not re.fullmatch(r'(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.0+)?', text):
        raise ValueError(f'Invalid count: {value!r}')
    try:
        number = Decimal(text.replace(',', ''))
    except InvalidOperation:
        raise ValueError(f'Invalid count: {value!r}') from None
    if not number.is_finite() or number < 0 or number > 2_147_483_647:
        raise ValueError(f'Count outside storage range: {value!r}')
    return int(number)


@dataclass
class Report:
    accepted: int = 0
    rejected: int = 0
    duplicates: int = 0
    missing: int = 0
    inserted: int = 0
    warnings: list[str] = field(default_factory=list)
    districts_covered: list[str] = field(default_factory=list)
    metrics_extracted: list[str] = field(default_factory=list)


def parse_rows(rows, config: dict) -> tuple[list[dict], Report]:
    """Rows carry Excel row numbers; reject all occurrences of duplicate identities."""
    report = Report()
    candidates = []
    seen = Counter()
    for row_number, row in rows:
        if not row:
            continue
        label = str(row[0] or '').strip()
        if not label.upper().endswith(' DISTRICT'):
            continue
        try:
            slug = resolve_district(label, config)
            seen[slug] += 1
            value = parse_count(row[2])
            if value is None:
                report.missing += 1
                report.warnings.append(f'Row {row_number}: missing population for {slug}')
                continue
            if value == 0:
                raise ValueError('District population must be positive')
            if value > config['warning_max']:
                report.warnings.append(f'Row {row_number}: population exceeds configured normal bound')
            candidates.append(dict(slug=slug, name=config['districts'][slug], value=value,
                metric_key=config['metric_labels']['POPULATION-2023'], unit=config['unit'],
                period=config['period'], source_id=config['source_id'],
                locator=f"Table 1; {config['sheet']}!A{row_number}:C{row_number}; "
                        f"source row={label!r}, population={row[2]!r}; "
                        f"extraction=openpyxl; parser={config['parser_version']}; sha256={config['sha256']}"))
        except (ValueError, IndexError) as exc:
            report.rejected += 1
            report.warnings.append(f'Row {row_number}: {exc}')
    accepted = [item for item in candidates if seen[item['slug']] == 1]
    report.duplicates = sum(count for count in seen.values() if count > 1)
    report.rejected += len(candidates) - len(accepted)
    for slug, count in seen.items():
        if count > 1:
            report.warnings.append(f'Duplicate district rows: {slug} ({count}); all withheld')
    report.accepted = len(accepted)
    report.districts_covered = sorted(item['slug'] for item in accepted)
    report.metrics_extracted = sorted({item['metric_key'] for item in accepted})
    absent = sorted(set(config['districts']) - set(report.districts_covered))
    if absent:
        report.warnings.append('Districts without accepted values: ' + ', '.join(absent))
    return accepted, report


def extract(path: Path, config: dict) -> tuple[list[dict], Report]:
    if config['period'] != '2023' or config['unit'] != 'count' or config['metric_labels'] != {'POPULATION-2023': 'population_2023'}:
        raise ValueError('Unsupported PBS metric, period, or unit contract')
    # A different edition requires an explicit reviewed configuration update.
    if sha256(path.read_bytes()).hexdigest() != config['sha256']:
        raise ValueError('Source checksum differs from the registered PBS Table 1 edition')
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sheet = workbook[config['sheet']]
        if (sheet['C2'].value != 'POPULATION-2023' or sheet['C3'].value != 'ALL SEXES'
                or 'CENSUS-2023, PUNJAB' not in str(sheet['A1'].value)
                or sheet['A2'].value != 'NAME OF ADMINISTRATIVE UNIT'):
            raise ValueError('Unexpected PBS Table 1 header layout')
        return parse_rows(enumerate(sheet.iter_rows(min_row=5, values_only=True), 5), config)
    finally:
        workbook.close()


def load_observations(session, observations: list[dict], source: dict, config: dict, report: Report):
    """Use existing uniqueness constraints; never overwrite reviewed data."""
    source_id = config['source_id']
    if source['period'] != config['period'] or not re.fullmatch(r'\d{4}', config['period']):
        raise ValueError('Invalid or conflicting reporting period')
    session.execute(insert(Source).values(**{**source, 'id': source_id, 'sha256': config['sha256']})
                    .on_conflict_do_nothing(index_elements=['id']))
    registered = session.get(Source, source_id)
    if any(getattr(registered, key) != source[key] for key in ('publisher', 'title', 'period', 'url')):
        raise ValueError('Registered source metadata conflicts with the configured source')
    if registered.sha256 not in (None, config['sha256']):
        raise ValueError('Registered source checksum conflicts with this edition')
    if registered.sha256 is None:
        registered.sha256 = config['sha256']
    for item in observations:
        population = item.get('population_2023', item['value'])
        session.execute(insert(District).values(slug=item['slug'], name=item['name'], population_2023=population)
                        .on_conflict_do_nothing(index_elements=['slug']))
        district = session.scalar(select(District).where(District.slug == item['slug']))
        if district.population_2023 != population or district.name != item['name']:
            report.accepted -= 1
            if item['slug'] in report.districts_covered:
                report.districts_covered.remove(item['slug'])
            report.rejected += 1
            report.warnings.append(f"{item['slug']}: conflicts with existing district; preserved")
            continue
        values = {key: item[key] for key in ('value', 'metric_key', 'unit', 'period', 'source_id', 'locator')}
        result = session.execute(insert(Observation).values(**values, district_id=district.id,
            category=item.get('category', 'context'), definition_version='prototype-v1', geography_level='district',
            population_basis=('Census 2023 district population' if item.get('category', 'context') == 'context' else None), quality_flags=[])
            .on_conflict_do_nothing(constraint='uq_observation_identity').returning(Observation.id)).scalar_one_or_none()
        if result is not None:
            report.inserted += 1
        else:
            existing = session.scalar(select(Observation).where(Observation.district_id == district.id,
                Observation.metric_key == item['metric_key'], Observation.period == item['period']))
            report.duplicates += 1
            if existing.value != Decimal(str(item['value'])) or existing.unit != item['unit'] or existing.source_id != source_id:
                report.accepted -= 1
                if item['slug'] in report.districts_covered:
                    report.districts_covered.remove(item['slug'])
                report.rejected += 1
                report.warnings.append(f"{item['slug']}: observation conflict; existing record preserved")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', choices=['pbs_table_1'], default='pbs_table_1')
    parser.add_argument('--dry-run', action='store_true', help='Validate without accessing PostgreSQL')
    parser.add_argument('--report', type=Path)
    args = parser.parse_args()
    config = json.loads(CONFIG.read_text())
    source = json.loads((ROOT / 'data/curated/districts.json').read_text())['sources'][args.source]
    try:
        observations, report = extract(ROOT / source['local_file'], config)
        if not args.dry_run:
            from app.config import Settings
            settings = Settings.from_env()
            if not settings.database_url:
                raise ValueError('GPIS_DATABASE_URL is required; use --dry-run for validation only')
            with session_factory(settings.database_url).begin() as session:
                load_observations(session, observations, source, config, report)
    except (ValueError, KeyError, OSError) as exc:
        parser.exit(2, f'Ingestion failed: {exc}\n')
    output = json.dumps(asdict(report), indent=2)
    if args.report:
        args.report.write_text(output + '\n', encoding='utf-8')
    print(output)
    return 1 if report.rejected or report.missing else 0


if __name__ == '__main__':
    raise SystemExit(main())
